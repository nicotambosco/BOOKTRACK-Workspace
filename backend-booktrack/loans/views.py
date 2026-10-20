from django.http import HttpResponse
from rest_framework import mixins, permissions, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from openpyxl import Workbook
from fpdf import FPDF
from .models import Loan, Reservation
from .serializers import LoanSerializer, ReservationSerializer, PrioritySerializer
from .permissions import IsLibrarian
from . import services


class LoanViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = LoanSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        services.refresh_states()
        queryset = Loan.objects.select_related('estudiante', 'libro').all()
        if self.request.user.categoria != 'bibliotecario':
            queryset = queryset.filter(estudiante=self.request.user)
        state = self.request.query_params.get('estado')
        if state:
            if state not in Loan.Estado.values:
                raise ValidationError({'estado': 'Estado de préstamo inválido.'})
            queryset = queryset.filter(estado=state)
        return queryset.order_by('-fecha_inicio', '-id')

    @action(detail=True, methods=['patch'], permission_classes=[IsLibrarian])
    def devolver(self, request, pk=None):
        return Response(self.get_serializer(services.return_loan(self.get_object())).data)

    @action(detail=True, methods=['patch'], url_path='solicitar-extension')
    def solicitar_extension(self, request, pk=None):
        return Response(self.get_serializer(services.request_extension(self.get_object())).data)

    @action(detail=True, methods=['patch'], url_path='aprobar-extension', permission_classes=[IsLibrarian])
    def aprobar_extension(self, request, pk=None):
        return Response(self.get_serializer(services.resolve_extension(self.get_object(), True)).data)

    @action(detail=True, methods=['patch'], url_path='denegar-extension', permission_classes=[IsLibrarian])
    def denegar_extension(self, request, pk=None):
        return Response(self.get_serializer(services.resolve_extension(self.get_object(), False)).data)

    @action(detail=False, methods=['get'], url_path='export/pdf')
    def export_pdf(self, request):
        loans = self.get_queryset()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font('Helvetica', size=14)
        pdf.cell(0, 10, 'Historial de prestamos', ln=True)
        pdf.ln(5)
        pdf.set_font('Helvetica', 'B', size=9)
        widths = [35, 40, 18, 20, 25, 25, 27]
        for width, label in zip(widths, ['Estudiante', 'Libro', 'Tipo', 'Estado', 'Inicio', 'Vencimiento', 'Devolucion']):
            pdf.cell(width, 7, label)
        pdf.ln()
        pdf.set_font('Helvetica', size=8)
        for loan in loans:
            cells = [(loan.estudiante.nombre_apellido or str(loan.estudiante))[:18], loan.libro.titulo[:20],
                     loan.tipo_prestamo[:10], loan.estado, str(loan.fecha_inicio),
                     str(loan.fecha_vencimiento or '-'), str(loan.fecha_devolucion or '-')]
            for width, text in zip(widths, cells):
                pdf.cell(width, 7, text.encode('latin-1', 'replace').decode('latin-1'))
            pdf.ln()
        response = HttpResponse(bytes(pdf.output()), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="historial-prestamos.pdf"'
        return response

    @action(detail=False, methods=['get'], url_path='export/excel')
    def export_excel(self, request):
        loans = self.get_queryset()
        wb = Workbook()
        ws = wb.active
        ws.title = 'Prestamos'
        ws.append(['Estudiante', 'Libro', 'Tipo', 'Estado', 'Fecha inicio', 'Fecha vencimiento', 'Fecha devolución'])
        for l in loans:
            nombre_estudiante = l.estudiante.nombre_apellido or str(l.estudiante)
            nombre_libro = l.libro.titulo if hasattr(l.libro, 'titulo') else str(l.libro)
            fecha_fin = str(l.fecha_vencimiento) if l.fecha_vencimiento else ''
            ws.append([nombre_estudiante, nombre_libro, l.tipo_prestamo, l.estado, str(l.fecha_inicio), fecha_fin, str(l.fecha_devolucion or "")])
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="historial-prestamos.xlsx"'
        wb.save(response)
        return response


class ReservationViewSet(mixins.CreateModelMixin, viewsets.ReadOnlyModelViewSet):
    serializer_class = ReservationSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        services.refresh_states()
        queryset = Reservation.objects.select_related('estudiante', 'libro').all()
        if self.request.user.categoria != 'bibliotecario':
            queryset = queryset.filter(estudiante=self.request.user)
        state = self.request.query_params.get('estado')
        if state:
            if state not in Reservation.Estado.values:
                raise ValidationError({'estado': 'Estado de reserva inválido.'})
            queryset = queryset.filter(estado=state)
        return queryset

    def perform_create(self, serializer):
        serializer.instance = services.create_reservation(estudiante=self.request.user, **serializer.validated_data)

    @action(detail=True, methods=['patch'])
    def cancelar(self, request, pk=None):
        return Response(self.get_serializer(services.cancel_reservation(self.get_object(), request.user)).data)

    @action(detail=True, methods=['patch'], permission_classes=[IsLibrarian])
    def retirar(self, request, pk=None):
        return Response(self.get_serializer(services.pick_up(self.get_object(), request.user)).data)

    @action(detail=True, methods=['patch'], permission_classes=[IsLibrarian])
    def priorizar(self, request, pk=None):
        data = PrioritySerializer(data=request.data)
        data.is_valid(raise_exception=True)
        reservation = services.prioritize_reservation(self.get_object(), data.validated_data['prioridad'])
        return Response(self.get_serializer(reservation).data)
